import json,copy,hashlib,subprocess,tempfile,shutil,math
from pathlib import Path
from datetime import datetime,timezone,timedelta
import openpyxl
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.utils import get_column_letter
p=Path("Live ТБ/Live ТБ.xlsx")
original_hash=hashlib.sha256(p.read_bytes()).hexdigest()
w=openpyxl.load_workbook(p)
cached=openpyxl.load_workbook(p,data_only=True)
d=json.loads(Path("diagnostics/diary-update-20261016.json").read_text())
a=json.loads(Path("reports/Данные расчёта live 15 октября 2026.json").read_text())
s=w["Реестр"]
old=[[c.value for c in row] for row in s.iter_rows(min_row=6,max_row=146,max_col=33)]
assert len(old)==141 and all(isinstance(r[0],(int,float)) for r in old)
for u in d["history_updates"]:
 row=5+u["n"]
 assert str(s.cell(row,2).value)==u["id"]
 for col,val in {20:u["final"],21:u["result"],22:u["profit"],23:"Рассчитан",24:"Результат подтверждён Events API 16.10.2026; расчёт за основное время по scores.2, без пенальти.",32:"BetsAPI /v1/event/view; 16.10.2026; scores.2"}.items(): s.cell(row,col,val)
headers={34:"ID снимка",35:"Справочный ТБ, матч",36:"Справочный коэффициент",37:"Результат справочного ТБ",38:"Справочный итог, ед.",39:"Голов после снимка",40:"Тип записи"}
for col,val in headers.items(): s.cell(5,col,val)
seen=set()
for i,(r,sel) in enumerate(zip(d["new_rows"],a["selections"]),147):
 assert r["id"]==sel["event_id"] and r["id"] not in seen
 seen.add(r["id"])
 for col in range(1,41):
  s.cell(i,col)._style=copy.copy(s.cell(146,min(col,33))._style)
  s.cell(i,col).value=None
 dt=datetime.fromisoformat(sel["selected_at_utc"].replace("Z","+00:00")).astimezone(timezone(timedelta(hours=3))).strftime("%d.%m.%Y %H:%M:%S")
 vals={1:r["n"],2:r["id"],3:dt,4:r["league"],5:r["match"],6:r["minute"],7:r["score"],8:"Опубликованный спортивный отбор",9:"Справочная линия в колонках AI:AL",16:"Только сохранённый снимок; признаки не реконструированы",17:"Опубликован",18:sel["rules_version"],19:"Технический реестр опубликованных Live; Events API",20:r["final"],21:r["result"],23:"Рассчитан справочно" if r["price"] else "Нет справочной линии",24:"Не фактическая ставка пользователя. Справочная цена на момент снимка; доходность условная.",25:"Спортивный отбор",26:sel["snapshot_id"]+"-"+r["id"],27:"Матч",29:dt,30:dt,31:"Нет",32:"BetsAPI /v1/event/view; scores.2; отчёт 15.10.2026",33:sel["rules_version"],34:sel["snapshot_id"],35:r["line"],36:r["price"],37:r["result"],38:r["profit"],39:r["extra"],40:"Справочная линия"}
 for col,val in vals.items():
  if val is not None:s.cell(i,col,val)
for row in range(6,147):s.cell(row,40,"Исторический прогноз")
assert len(seen)==72
# Historical core remains intact except explicitly settled rows.
allowed={20,21,22,23,24,32}
changed_numbers={u["n"] for u in d["history_updates"]}
for idx,vals in enumerate(old,6):
 for col,val in enumerate(vals,1):
  if idx-5 in changed_numbers and col in allowed:continue
  assert s.cell(idx,col).value==val,(idx,col)
s.auto_filter.ref="A5:AN218"
s.freeze_panes="F6"
for col in range(34,41):s.column_dimensions[get_column_letter(col)].width=24
s.print_options.horizontalCentered=True
# Preserve previous summary as an explicitly dated archive with its cached values.
archive=w["Сводка"]
archive.title="Сводка до 16.10"
for row in archive:
 for c in row:
  if c.data_type=="f":c.value=cached["Сводка"][c.coordinate].value
archive["A3"]="Архив прежней сводки до обновления 16.10.2026. Актуальные итоги на листе Сводка."
summary=w.create_sheet("Сводка",0)
summary.append(["Live ТБ: обновлено 16.10.2026"])
summary.append(["213 записей / 170 уникальных матчей. Исторические прогнозы и справочные линии разделены."])
summary.append(["Показатель","Значение"])
summary.append(["Всего записей",213])
summary.append(["Уникальных матчей",170])
summary.append(["Исторических прогнозов",141])
summary.append(["Спортивных выборов 15 октября",72])
summary.append(["Из них со справочной ценой",68])
summary.append(["Без справочной линии",4])
summary.append(["Досчитано исторических строк",14])
summary.append([])
summary.append(["Группа","Строк","Уникальных матчей","Итог, ед.","ROI","Примечание"])
for g in d["analysis"]:
 if g["dim"]=="contour" and g["cohort"] in ["История 13–14: все строки","15 октября: справочная линия"]:
  summary.append([g["k"],g["n"],g["u"],g["p"],g["roi"]/100,"1 условная единица на строку; это не личный денежный результат"])
  summary.cell(summary.max_row,5).number_format="0.00%"
summary.append([])
summary.append(["Спортивные результаты 15 октября","Матчей","Доля"])
for label,num in [("Хотя бы один новый гол",68),("Хотя бы два новых гола",53)]:
 summary.append([label,num,num/72]);summary.cell(summary.max_row,3).number_format="0.00%"
summary.append([])
summary.append(["Ограничения"])
for note in ["Всего три дня и несколько версий отбора. Найденные плюсы являются гипотезами, не подтверждением будущей доходности.",
"Старые строки учитывают несколько линий и повторные входы. Для контроля на листе Анализ есть первый прогноз на весь матч на каждое событие.",
"В новых строках коэффициент и расчёт только в справочных колонках AI:AL. В исторический финансовый итог V они не включены.",
"Старые подтверждённые результаты сохранены по дневнику; повторно через API проверялись 12 ранее незакрытых матчей.",
"Первый тайм нельзя смешивать с тоталом матча. Персональные ставки и суммы в эту книгу не переносились.",
"Правила и код сканера этим обновлением не менялись. Архив Правила может отражать прежнюю редакцию; действующий канон находится в репозитории."]:
 summary.append([note])
an=w.create_sheet("Анализ")
an.append(["Исследовательские срезы на 16.10.2026. Все группы, включая отрицательные; ROI по единичной ставке."])
an.append(["Когорта","Признак","Группа","Строк","Уникальных матчей","Итог, ед.","ROI","Размер выборки"])
labels={"contour":"Контур","period":"Период","minute":"Минута","score":"Голы на снимке","league":"Лига","price":"Коэффициент","needed":"Доп. голов для полного выигрыша","category":"Тип команд"}
for g in d["analysis"]:
 an.append([g["cohort"],labels[g["dim"]],g["k"],g["n"],g["u"],g["p"],g["roi"]/100,"Очень малая" if g["u"]<10 else "Недостаточно для подтверждения"])
 an.cell(an.max_row,7).number_format="0.00%"
an.freeze_panes="D3";an.auto_filter.ref=f"A2:H{an.max_row}"
for sheet in [summary,an]:
 for row in sheet:
  for c in row:
   c.font=Font(name="Times New Roman",size=11)
   c.alignment=Alignment(vertical="top",wrap_text=True)
 for c in sheet[1]:c.font=Font(name="Times New Roman",size=11,bold=True)
 sheet.sheet_view.showGridLines=False
 sheet.sheet_properties.pageSetUpPr.fitToPage=True
 sheet.page_setup.orientation="landscape";sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
 sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
for col,width in {"A":52,"B":32,"C":38,"D":18,"E":18,"F":70}.items():summary.column_dimensions[col].width=width
for row in summary:
 if row[0].value and len(str(row[0].value))>80:summary.row_dimensions[row[0].row].height=72
for col,width in {"A":44,"B":38,"C":52,"D":12,"E":16,"F":18,"G":15,"H":37}.items():an.column_dimensions[col].width=width
for r in range(1,an.max_row+1):an.row_dimensions[r].height=34
for row in [5]:
 for c in s[row]:
  c.font=Font(name="Times New Roman",size=11,bold=True);c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
s.row_dimensions[5].height=48
for row in range(147,219):
 s.row_dimensions[row].height=65
 for c in s[row]:
  c.font=Font(name="Times New Roman",size=11);c.alignment=Alignment(vertical="top",wrap_text=True)
# Retain caches as values for any preexisting formula outside the regenerated summary.
for sh in w:
 if sh.title in ["Сводка","Анализ","Сводка до 16.10"]:continue
 for row in sh:
  for c in row:
   if c.data_type=="f" and sh.title in cached.sheetnames:
    val=cached[sh.title][c.coordinate].value
    assert val is not None,("uncached formula",sh.title,c.coordinate)
    c.value=val
w.save(p)
v=openpyxl.load_workbook(p,data_only=True)
rows=[r for r in v["Реестр"].iter_rows(min_row=6,values_only=True) if isinstance(r[0],(int,float)) and r[4]]
assert len(rows)==213
assert len({str(r[1] or r[4]) for r in rows})==170
assert sum(r[39]=="Справочная линия" for r in rows)==72
assert abs(sum(r[37] or 0 for r in rows)+6.475)<1e-8
assert all(r[21] is None for r in rows[-72:])
assert all(r[21] is not None for r in rows[:141])
report=Path("reports/live-diary-audit-20261016.json")
report.write_text(json.dumps({"updated_at_utc":datetime.now(timezone.utc).isoformat(),"records":213,"unique_matches":170,"historical_records":141,"historical_profit":sum(r[21] for r in rows[:141]),"new_records":72,"reference_priced":68,"reference_profit":-6.475,"settled_rows":d["history_updates"],"pending_api_responses":d["pending"],"groups":d["analysis"],"limitations":["Exploratory, three days only","Mixed rule versions, separate cohorts","Old settled results retained from diary, not all reverified","No personal betting diary data"]},ensure_ascii=False,indent=2))
# Publish only the workbook and public audit report onto the latest main.
subprocess.run(["git","fetch","origin","main"],check=True,capture_output=True)
tmp=Path(tempfile.mkdtemp())/"main"
subprocess.run(["git","worktree","add","--detach",str(tmp),"origin/main"],check=True,capture_output=True)
assert hashlib.sha256((tmp/p).read_bytes()).hexdigest()==original_hash,"Main workbook changed concurrently; refusing overwrite"
for path in [p,report]:
 (tmp/path).parent.mkdir(parents=True,exist_ok=True)
 shutil.copy2(path,tmp/path)
def git(*args):return subprocess.run(["git","-C",str(tmp),*args],check=True,capture_output=True,text=True).stdout.strip()
git("config","user.name","github-actions[bot]")
git("config","user.email","41898282+github-actions[bot]@users.noreply.github.com")
git("add",str(p),str(report))
git("commit","-m","Update Live diary: 72 selections, settle 14 historical rows, add exploratory analysis")
git("push","origin","HEAD:main")
print("UPDATE_RESULT="+json.dumps({"commit":git("rev-parse","HEAD"),"records":len(rows),"unique":170,"historical_profit":sum(r[21] for r in rows[:141]),"reference_profit":sum(r[37] or 0 for r in rows)},ensure_ascii=False))
